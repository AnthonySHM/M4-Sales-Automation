"""
Build the MSS SOW master template Word document.
Run: python scripts/build_master_template.py
Output: Desktop/MSS_SOW_Master_Template_DRAFT.docx
"""
from docx import Document
from docx.shared import Pt, RGBColor
import os

doc = Document()


def marker(d, text):
    """Red bold paragraph for block markers."""
    para = d.add_paragraph()
    run = para.add_run(text)
    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    run.font.bold = True
    run.font.size = Pt(9)
    return para


def h(d, text, level=1):
    d.add_heading(text, level=level)


def p(d, text):
    d.add_paragraph(text)


def fee_table(d, tier_name, sku, has_isensor_footnote=True):
    """Add the three fee tables (POC, Term, Remainder) for a tier."""

    if tier_name == "Enterprise":
        rows_data = [
            ("Mantix4 Enterprise", sku),
            ("Mantix4 Additional Data Storage Service2", "M4ADSS2003"),
            ("SentinelOne Singularity XDR Licensing", ""),
            ("Other?", ""),
        ]
    else:
        rows_data = [
            (f"Mantix4 {tier_name}", sku),
            ("Mantix4 SIEM Service", "M4SIEM2001"),
            ("Mantix4 Firewall Rule Compliance Monitoring Module", "M4FWRM2002"),
            ("Mantix4 Additional Data Storage Service2", "M4ADSS2003"),
            ("Mantix4 SecurityBox Add-on - Endpoint Patching Assessment", "M4EPAS2004"),
            ("SentinelOne Singularity XDR Licensing", ""),
            ("Other?", ""),
        ]

    headers = ["Service", "SKU", "Frequency", "Quantity*", "Fee per Service1", "Subtotal Monthly Fees"]

    # TABLE 1: Services overview
    t1 = d.add_table(rows=2, cols=2)
    t1.style = "Table Grid"
    t1.rows[0].cells[0].text = "SERVICES"
    t1.rows[0].cells[1].text = "Mantix4 Managed Cyber Security Services as described below"
    t1.rows[1].cells[0].text = "Fees for additional services will be provided in separate Quote"
    d.add_paragraph()

    # TABLE 2: POC fee header
    t2 = d.add_table(rows=1, cols=2)
    t2.style = "Table Grid"
    t2.rows[0].cells[0].text = "FEES (Monthly) for Initial Two (2) Months of Services (Proof of Concept Period)"
    t2.rows[0].cells[1].text = "{{POC_FEE_MONTHLY}}"
    d.add_paragraph()

    def itemized_table(d, rows_data, headers):
        n = len(rows_data)
        t = d.add_table(rows=n + 2, cols=6)
        t.style = "Table Grid"
        for i, h_text in enumerate(headers):
            t.rows[0].cells[i].text = h_text
        for r, (svc, s) in enumerate(rows_data):
            t.rows[r + 1].cells[0].text = svc
            t.rows[r + 1].cells[1].text = s
            t.rows[r + 1].cells[2].text = "Per Endpoint/ Per Month"
            t.rows[r + 1].cells[3].text = ""
            t.rows[r + 1].cells[4].text = "$ 0.00"
            t.rows[r + 1].cells[5].text = "$ 0.00"
        t.rows[n + 1].cells[0].text = "Total Monthly Fees1 (Recurring)"
        t.rows[n + 1].cells[5].text = "$ 0.00"
        d.add_paragraph()

    # TABLE 3: POC itemized
    itemized_table(d, rows_data, headers)

    # TABLE 4: Term fee header
    t4 = d.add_table(rows=1, cols=2)
    t4.style = "Table Grid"
    t4.rows[0].cells[0].text = "FEES (Monthly) for Term of Services"
    t4.rows[0].cells[1].text = "{{TERM_FEE_MONTHLY}}"
    d.add_paragraph()

    # TABLE 5: Term itemized
    itemized_table(d, rows_data, headers)

    # TABLE 6: Remainder fee header
    t6 = d.add_table(rows=1, cols=2)
    t6.style = "Table Grid"
    t6.rows[0].cells[0].text = "FEES (Monthly) for Remainder of Term of Services"
    t6.rows[0].cells[1].text = "{{REMAINDER_FEE_MONTHLY}}"
    d.add_paragraph()

    # TABLE 7: Footnotes
    t7 = d.add_table(rows=3, cols=2)
    t7.style = "Table Grid"
    t7.rows[0].cells[0].text = "*Minimum Endpoint Count: {{MIN_ENDPOINT_COUNT}}"
    t7.rows[1].cells[0].text = "1 Fee schedule is in U.S. Dollars ($USD)."
    t7.rows[2].cells[0].text = "2 Mantix4 Additional Data Storage Service will be custom quoted to meet specific Compliance requirements."
    if has_isensor_footnote:
        t7.rows[2].cells[1].text = "3 M4 Hardware iSensor is provided as a Service (HaaS), remains property of Mantix4, and will be returned with termination of contracted services."
    d.add_paragraph()


# ============================================================
# COVER PAGE
# ============================================================
doc.add_heading("Managed Security Services", 0)
doc.add_heading("Statement of Work", 2)
p(doc, "Prepared for")
p(doc, "{{CUSTOMER_LOGO}}")
p(doc, "{{CUSTOMER_NAME}}")
p(doc, "{{CUSTOMER_CONTACT_NAME}}, {{CUSTOMER_CONTACT_TITLE}}")
p(doc, "")
p(doc, "Prepared by")
p(doc, "Account Manager: {{ACCOUNT_MANAGER_NAME}}")
p(doc, "Solutions Architect: Chris Dodunski / Chuck Taylor")
p(doc, "Date: {{SOW_DATE}}")
p(doc, "Quote #: {{QUOTE_NUMBER}}")
p(doc, "Revision #: {{REVISION_NUMBER}}")
p(doc, "")

marker(doc, "[[BLOCK:partner_logo]]")
p(doc, "Partner Logo")
marker(doc, "[[/BLOCK:partner_logo]]")

doc.add_page_break()

# ============================================================
# DEFINITION OF PARTIES
# ============================================================
h(doc, "Definition of Parties", 1)
p(doc, (
    'This Statement of Work ("SOW") is between {{CUSTOMER_NAME}} ("Customer") and Mantix4, LLC ("Company"), '
    '(each a "Party", and collectively, the "Parties"), for the purpose of performing the tasks as defined herein '
    "on an ongoing basis for the specified period of the Term."
))
p(doc, (
    'This Managed Security Services SOW is effective by and between Customer and Company pursuant to the signed '
    'Master Services Agreement ("Agreement") dated {{MSA_DATE}}, by and between the Parties. To the extent there '
    "are any conflicts or inconsistencies between the Agreement and this SOW, the provisions of this SOW shall "
    "control and prevail, but only with respect to this SOW. Capitalized terms used in this SOW that are not "
    "otherwise defined below shall have the meanings in the Agreement or schedules, appendices or exhibits to the Agreement."
))

# ============================================================
# NON-DISCLOSURE
# ============================================================
h(doc, "Non-Disclosure", 1)
p(doc, (
    "This proposal and its contents are confidential Company information and methodologies and subject to the "
    "non-disclosure provisions set forth in the Agreement. In consideration of the receipt of this document, "
    "Customer agrees not to reproduce, distribute or make this information available in any manner to persons "
    "outside the Customer directly responsible for evaluation of its contents."
))

# ============================================================
# SERVICES AND FEES
# ============================================================
h(doc, "Services and Fees", 1)
h(doc, "Mantix4 Managed Security Services", 2)

# ---- SERVICE DESCRIPTIONS (mutually exclusive, one per tier) ----

marker(doc, "[[BLOCK:service_desc_elite]]")
h(doc, "Mantix4 Integrated MDR \u2013 Elite (Proof of Concept (POC))", 3)
p(doc, (
    "M4 SecurityBox AI/ML powers Mantix4\u2019s full Managed Detection and Response (MDR) and connects the Customer\u2019s "
    "own AI to a full M4 security team. Agentic enhancement of monitoring and analysis, leading with Network and ingested "
    "Cloud Services, Identity, Endpoint, Access, and Exposure telemetry (cloud environment and on-prem), results in one "
    "correlated view with evidence behind every finding, and enhanced detection, response, and recovery from the most "
    "advanced cyber-attacks. M4 SecurityBox automated response is available and will be tuned to Customer\u2019s needs."
))
p(doc, (
    "Bring Your Own AI: Mantix4 brings the Agentic security team. The Customer\u2019s AI connects over an authenticated, "
    "open-standard interface and receives capability, not credentials. No keys to sensors, consoles, identity tenants, "
    "or databases are handed over, and access is revocable in a single action."
))
p(doc, (
    "Continuous, proactive risk reduction: Mantix4\u2019s Agentic Continuous Threat Exposure Management (CTEM) hardens "
    "the environment ahead of attack, while reduced Mean Time to Detect, Respond, and Fix (MTTD / MTTR / MTTF) limits "
    "the damage of what does get through."
))
p(doc, "Coverage: Network, Cloud Services and Identity, Endpoints, Access, and Exposure (cloud environments and on-prem).")
p(doc, "Designed for organizations who have a hybrid physical - cloud network and cloud services, and their own AI.")
marker(doc, "[[/BLOCK:service_desc_elite]]")

marker(doc, "[[BLOCK:service_desc_enterprise]]")
h(doc, "Mantix4 Enterprise (MDR Enhancement) (Proof of Concept (POC))", 3)
p(doc, (
    "M4 SecurityBox AI/ML supercharges Managed Detection and Response (MDR) of the existing SIEM/SOC, and connects "
    "the Customer\u2019s own AI to a full M4 security team. Agentic enhancement of monitoring and analysis, leading "
    "with Network and ingested Cloud Services, Identity, Endpoint, Access, and Exposure telemetry (cloud environment "
    "and on-prem), results in one correlated view with evidence behind every finding, and enhanced detection, response, "
    "and recovery from the most advanced cyber-attacks. M4 SecurityBox automated response is available and will be "
    "tuned to Customer\u2019s needs."
))
p(doc, (
    "Bring Your Own AI: Mantix4 brings the Agentic security team. The Customer\u2019s AI connects over an authenticated, "
    "open-standard interface and receives capability, not credentials. No keys to sensors, consoles, identity tenants, "
    "or databases are handed over, and access is revocable in a single action."
))
p(doc, (
    "Continuous, proactive risk reduction: Mantix4\u2019s Agentic Continuous Threat Exposure Management (CTEM) hardens "
    "the environment ahead of attack, while reduced Mean Time to Detect, Respond, and Fix (MTTD / MTTR / MTTF) limits "
    "the damage of what does get through."
))
p(doc, "Coverage: Network, Cloud Services and Identity, Endpoints, Access, and Exposure (cloud environments and on-prem).")
p(doc, "Designed for organizations who have a hybrid physical - cloud network and cloud services, their own SIEM/SOC, and their own AI.")
marker(doc, "[[/BLOCK:service_desc_enterprise]]")

marker(doc, "[[BLOCK:service_desc_premium]]")
h(doc, "Mantix4 Integrated MDR \u2013 Premium (Proof of Concept (POC))", 3)
p(doc, (
    "M4 SecurityBox AI/ML powers Mantix4\u2019s full Managed Detection and Response (MDR) and connects the Customer\u2019s "
    "own AI to a full M4 security team. Agentic enhancement of monitoring and analysis, leading with Network and ingested "
    "Endpoint, Access, and Exposure telemetry (cloud environment and on-prem), results in one correlated view with "
    "evidence behind every finding, and enhanced detection, response, and recovery from the most advanced cyber-attacks. "
    "M4 SecurityBox automated response is available and will be tuned to Customer\u2019s needs."
))
p(doc, (
    "Bring Your Own AI: Mantix4 brings the Agentic security team. The Customer\u2019s AI connects over an authenticated, "
    "open-standard interface and receives capability, not credentials. No keys to sensors, consoles, identity tenants, "
    "or databases are handed over, and access is revocable in a single action."
))
p(doc, (
    "Continuous, proactive risk reduction: Mantix4\u2019s Agentic Continuous Threat Exposure Management (CTEM) hardens "
    "the environment ahead of attack, while reduced Mean Time to Detect, Respond, and Fix (MTTD / MTTR / MTTF) limits "
    "the damage of what does get through."
))
p(doc, "Coverage: Network, Endpoints, Access, and Exposure (cloud environments and on-prem).")
p(doc, "Designed for organizations who have a physical or hybrid - physical network and no cloud services, and their own AI.")
marker(doc, "[[/BLOCK:service_desc_premium]]")

marker(doc, "[[BLOCK:service_desc_strategic]]")
h(doc, "Mantix4 Integrated MDR \u2013 Stratus (Proof of Concept (POC))", 3)
p(doc, (
    "M4 SecurityBox AI/ML powers Mantix4\u2019s full Managed Detection and Response (MDR) and connects the Customer\u2019s "
    "own AI to a full M4 security team. Agentic enhancement of monitoring and analysis, leading with Cloud Services "
    "and Identity telemetry, ingesting Endpoint, Access, and Exposure telemetry (cloud environment and on-prem), "
    "results in one correlated view with evidence behind every finding, and enhanced detection, response, and recovery "
    "from the most advanced cyber-attacks. M4 SecurityBox automated response is available and will be tuned to "
    "Customer\u2019s needs."
))
p(doc, (
    "Bring Your Own AI: Mantix4 brings the Agentic security team. The Customer\u2019s AI connects over an authenticated, "
    "open-standard interface and receives capability, not credentials. No keys to sensors, consoles, identity tenants, "
    "or databases are handed over, and access is revocable in a single action."
))
p(doc, (
    "Continuous, proactive risk reduction: Mantix4\u2019s Agentic Continuous Threat Exposure Management (CTEM) hardens "
    "the environment ahead of attack, while reduced Mean Time to Detect, Respond, and Fix (MTTD / MTTR / MTTF) limits "
    "the damage of what does get through."
))
p(doc, "Coverage: Cloud Services and Identity, Endpoints, Access, and Exposure (cloud environments and on-prem).")
p(doc, "Designed for organizations who work in the cloud and have no physical network, and their own AI.")
marker(doc, "[[/BLOCK:service_desc_strategic]]")

# ---- MDR SECTIONS (mutually exclusive, Enterprise has none) ----

marker(doc, "[[BLOCK:mdr_elite]]")
h(doc, "Mantix4 Managed Detection & Response (MDR)", 3)
p(doc, "24 x 7 Continuous Detection and Response Monitoring.")
p(doc, "Leads with M4 NDR (Network Traffic Analysis / Threat Hunting).")
p(doc, "Service integration and analysis of Cloud Services (Microsoft 365, Google Workspace, and more): identity/access, network traffic, security configuration, and related telemetry.")
p(doc, "Service integration and analysis of Endpoints (cloud environments such as Azure and AWS, and on-prem): EDR/XDR alerts/logs and Endpoint logs.")
p(doc, "SIEM Service for Compliance needs available for an additional fee.")
marker(doc, "[[/BLOCK:mdr_elite]]")

marker(doc, "[[BLOCK:mdr_premium]]")
h(doc, "Mantix4 Managed Detection & Response (MDR)", 3)
p(doc, "24 x 7 Continuous Detection and Response Monitoring.")
p(doc, "Leads with M4 NDR (Network Traffic Analysis / Threat Hunting).")
p(doc, "Service integration and analysis of Endpoints (cloud environments such as Azure and AWS, and on-prem): EDR/XDR alerts/logs and Endpoint logs.")
p(doc, "SIEM Service for Compliance needs available for an additional fee.")
marker(doc, "[[/BLOCK:mdr_premium]]")

marker(doc, "[[BLOCK:mdr_strategic]]")
h(doc, "Mantix4 Managed Detection & Response (MDR)", 3)
p(doc, "24 x 7 Continuous Detection and Response Monitoring.")
p(doc, "Service integration and analysis of Cloud Services (Microsoft 365, Google Workspace, and more): identity/access, network traffic, security configuration, and related telemetry.")
p(doc, "Service integration and analysis of Endpoints (cloud environments such as Azure and AWS, and on-prem): EDR/XDR alerts/logs and Endpoint logs.")
p(doc, "SIEM Service for Compliance needs available for an additional fee.")
marker(doc, "[[/BLOCK:mdr_strategic]]")

# ---- M4 SECURITYBOX (always present for all tiers) ----
h(doc, "M4 SecurityBox", 3)
p(doc, "Actionable Intelligence: Alert, Investigate, Hunt, and Respond.")
p(doc, "Highly Curated Security Event / Triage Alerts and Notifications that include Guided Remediation.")
p(doc, (
    "Correlated multi-source investigation with evidence and provenance behind every claim. Threat hunts initiated "
    "by the Customer\u2019s Security Team and AI run asynchronously across every altitude. Proposed response actions "
    "require human approval and are recorded end to end."
))
p(doc, "M4 SecurityBox Analyst Bench: Specialists your AI can put to work.")

marker(doc, "[[BLOCK:bench_ndr]]")
p(doc, (
    "Network Detection and Response (NDR): Mantix4\u2019s proactive AI/ML enhanced Network Traffic Analysis and "
    "Threat Hunting against the Customer\u2019s own baseline, surfacing the traffic that looks ordinary until "
    "the baseline says otherwise."
))
marker(doc, "[[/BLOCK:bench_ndr]]")

marker(doc, "[[BLOCK:bench_cloud_identity]]")
p(doc, (
    "Cloud and Identity: Microsoft 365, Entra, and Azure identity, access, and configuration behavior across "
    "the cloud services the Customer actually runs, where most attacks are identity events before they are anything else."
))
marker(doc, "[[/BLOCK:bench_cloud_identity]]")

p(doc, (
    "Endpoint (XDR/EDR): EDR/XDR telemetry alongside Windows event logs, providing the host context that separates "
    "a suspicious command line from a compromised machine. EDR/XDR software licensing available separately."
))
p(doc, (
    "Access (IAM): Who has what, who should not, and who quietly gained it: standing privilege, stale access, "
    "and the escalation paths that lead to critical assets."
))
p(doc, (
    "Exposure: What is installed, what version, what is unpatched, and what is reachable from outside: "
    "the Customer\u2019s standing weakness, inventoried continuously."
))
p(doc, "Agentic Continuous Threat Exposure Management (CTEM):")
p(doc, (
    "Continuous threat exposure management across everything the Analyst Bench observes. Real security posture "
    "measured over time and translated into business risk, weighing exposure against each asset\u2019s identity, "
    "access, and traffic context. Remediation prioritized in the order an attacker would exploit it, with detailed "
    "mitigation instructions."
))
p(doc, "Reporting and Analytics:")
p(doc, "On-demand C-Suite and Tactical reporting from one underlying picture, with Risk Management Strategy coaching of the Customer CISO/CIO/IT Leader.")
p(doc, "Data Storage and Onboarding:")
p(doc, "Mantix4 hosted one (1) year security event data storage on infrastructure dedicated to the Customer. Additional data storage available to meet Compliance requirements for an additional fee.2")

marker(doc, "[[BLOCK:isensor]]")
p(doc, "M4 SecurityBox / NDR iSensor(s) (Hardware / Virtual). Hardware iSensors are preferred and delivered as a Service (HaaS) for no additional fee.3")
marker(doc, "[[/BLOCK:isensor]]")

p(doc, "Mantix4 white glove technical onboarding support.")

# ---- ADD-ONS ----

marker(doc, "[[BLOCK:addon_siem]]")
h(doc, "Mantix4 Security Information and Event Management (SIEM) Service", 3)
p(doc, (
    "Supports meeting Compliance needs through automation of the collection, analysis, and reporting of security "
    "log data required by regulations like GDPR, HIPAA, PCI DSS, and SOX. Includes Mantix4 SIEM Endpoint software "
    "licensing and white glove technical onboarding support."
))
marker(doc, "[[/BLOCK:addon_siem]]")

marker(doc, "[[BLOCK:addon_firewall_compliance]]")
h(doc, "Mantix4 Firewall Rule Compliance Monitoring Module", 3)
p(doc, (
    "Includes continuous Firewall Rule Compliance Monitoring to help assess whether an organization is adhering "
    "to regulatory requirements, including internal policies and specific industry standards."
))
p(doc, "Reporting and Analytics including: critical notifications and Firewall Rule Compliance Assessment Reports with recommended remediation actions as necessary.")
marker(doc, "[[/BLOCK:addon_firewall_compliance]]")

marker(doc, "[[BLOCK:addon_storage]]")
h(doc, "Mantix4 Additional Data Storage Service (Compliance)", 3)
p(doc, "{{HOT_STORAGE_MONTHS}} Months Hot Storage; {{COLD_STORAGE_MONTHS}} Months Cold Storage")
marker(doc, "[[/BLOCK:addon_storage]]")

marker(doc, "[[BLOCK:addon_endpoint_patching]]")
h(doc, "Mantix4 SecurityBox Add-on - Endpoint Patching Assessment", 3)
p(doc, "Includes Monthly Assessment Reports for Microsoft Windows Systems, Third-Party software, and Browsers.")
marker(doc, "[[/BLOCK:addon_endpoint_patching]]")

h(doc, "Third Party Licensing Services & Fees", 2)

marker(doc, "[[BLOCK:addon_sentinelone]]")
h(doc, "SentinelOne Singularity XDR Licensing", 3)
p(doc, (
    "SentinelOne security platform employs AI enhanced endpoint protection against various threats, including "
    "malware, ransomware, zero-day and other advanced persistent threats (APTs). It detects and responds in "
    "real-time, ensuring that endpoints are safeguarded against both known and emerging threats."
))
p(doc, (
    "SentinelOne XDR to be installed on all endpoints. The platform, once a high degree of confidence is attained "
    "through MITRE ATT&CK\u00ae driven techniques, is designed to automatically respond to endpoint compromises."
))
p(doc, (
    "SentinelOne Extended Endpoint Protection Licensing (Anti-Ransomware) is a superior addition to Mantix4 "
    "Integrated MDR. When added, Mantix4 SIEM integrates SentinelOne\u2019s best in class Endpoint actions and "
    "alerts into its Network based Detection and Response providing full context visibility based proactive "
    "defense of Customer\u2019s IT network."
))
marker(doc, "[[/BLOCK:addon_sentinelone]]")

# ---- ONE-TIME SERVICES ----
h(doc, "One-Time and Ad Hoc Services", 2)
p(doc, "Any one-time or ad hoc services not specifically presented in this statement of work will be proposed separately with a scope of work.")
p(doc, "Prior to billing for the month, Company reserves the option to calculate the number of devices covered under the service and update the billing for the applicable month.")

# ---- SLO SECTION ----
marker(doc, "[[BLOCK:slo_section]]")
h(doc, "Mantix4 Service Level Objectives (SLO)", 2)
p(doc, "Response. Cases and standards are defined by the following table.")
slo_t = doc.add_table(rows=3, cols=4)
slo_t.style = "Table Grid"
slo_t.rows[0].cells[0].text = "Case Type"
slo_t.rows[0].cells[1].text = "Priority"
slo_t.rows[0].cells[2].text = "Response Time"
slo_t.rows[0].cells[3].text = "Resolution Target"
slo_t.rows[1].cells[0].text = "[fill from existing templates]"
slo_t.rows[2].cells[0].text = "[fill from existing templates]"
p(doc, "")
marker(doc, "[[/BLOCK:slo_section]]")

# ---- ORDER START DATE, TERM, POC, INVOICING ----
p(doc, "Order Start Date: {{ORDER_START_DATE}}")

h(doc, "Term", 2)
p(doc, (
    'Agreed Services begin on Order Start Date and continue until the earlier of three (3) years, or until '
    'terminated for convenience by either party with ninety (90) days written notice (the "Term"), or earlier '
    "in accordance with the Master Services Agreement. After the initial three (3) year period the Agreement "
    "will automatically renew for one (1) year terms until canceled by either party with ninety (90) days "
    "written notice."
))

h(doc, "Proof of Concept (POC)", 2)
p(doc, (
    "As presented in Monthly Services Fees above, the two-months of services under this POC {{SERVICE_TIER_FULL_NAME}} "
    "agreement are provided at no cost. If at any time Customer is not satisfied with these POC Services within "
    "the two-month POC Period plus fifteen (15) days after, notify Mantix4 Sales Account Manager and the Services "
    "will be cancelled and all tools expeditiously removed. If Customer wishes to retain the additional Services "
    "beyond POC Period, this agreement will continue in effect; non-cancellation is considered affirmation. "
    "Third party Licensing will be charged at regular price during the POC period."
))

h(doc, "Invoicing Terms", 2)
p(doc, (
    "All service charges will be accrued as agreed upon via execution of this SOW to be paid upon progression "
    "billing by calendar month. Onboarding Fees are waived. Billing begins on the first day of Services including "
    "the implementation time frame of the Services. Payment terms are in accordance with the Master Services Agreement."
))

p(doc, "IN WITNESS WHEREOF, the parties hereto have caused this Agreement to be duly executed set forth under their respective signatures.")

# ============================================================
# PRICING TABLES (one block per tier)
# ============================================================
doc.add_page_break()
h(doc, "PRICING TABLES", 1)
p(doc, "The engine selects exactly one of the four blocks below based on the chosen tier. All other three are deleted.")

marker(doc, "[[BLOCK:fees_table_elite]]")
h(doc, "Elite Pricing (M4ELIT1002)", 2)
fee_table(doc, "Elite", "M4ELIT1002", has_isensor_footnote=True)
marker(doc, "[[/BLOCK:fees_table_elite]]")

marker(doc, "[[BLOCK:fees_table_enterprise]]")
h(doc, "Enterprise Pricing (M4ENTN1001)", 2)
fee_table(doc, "Enterprise", "M4ENTN1001", has_isensor_footnote=True)
marker(doc, "[[/BLOCK:fees_table_enterprise]]")

marker(doc, "[[BLOCK:fees_table_premium]]")
h(doc, "Premium Pricing (M4PREM1003)", 2)
fee_table(doc, "Premium", "M4PREM1003", has_isensor_footnote=True)
marker(doc, "[[/BLOCK:fees_table_premium]]")

marker(doc, "[[BLOCK:fees_table_strategic]]")
h(doc, "Strategic / Stratus Pricing (M4STRAT1004)", 2)
fee_table(doc, "Stratus", "M4STRAT1004", has_isensor_footnote=False)
marker(doc, "[[/BLOCK:fees_table_strategic]]")

# ---- SIGNATURE ----
sig_t = doc.add_table(rows=4, cols=2)
sig_t.style = "Table Grid"
sig_t.rows[0].cells[0].text = "Customer Approval Authority"
sig_t.rows[0].cells[1].text = "Customer Approval Authority"
sig_t.rows[1].cells[0].text = "Name:"
sig_t.rows[1].cells[1].text = "Title:"
sig_t.rows[2].cells[0].text = "Signature:"
sig_t.rows[2].cells[1].text = "Date:"
sig_t.rows[3].cells[0].text = "Company Approval Authority"
sig_t.rows[3].cells[1].text = "Company Approval Authority"

out = os.path.join(os.path.expanduser("~"), "Desktop", "MSS_SOW_Master_Template_DRAFT.docx")
doc.save(out)
print(f"Saved: {out}")
